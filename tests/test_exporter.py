from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from monge_crawler.models import Post, FetchError
from monge_crawler.exporter import export


def test_export_creates_two_sheets_with_data(tmp_path):
    posts = [Post(post_id="1002", title="일반글 A", author="홍길동",
                  created_at=datetime(2025, 9, 12, 10, 30, 0),
                  body_text="본문", image_urls=["u1", "u2"],
                  image_paths=["p1", "p2"], url="https://article/1002")]
    errors = [FetchError(post_id="1003", url="https://article/1003", stage="article",
                         reason="timeout", at=datetime(2025, 9, 12, 10, 31, 0))]
    out = export(posts, errors, tmp_path / "out.xlsx")
    assert Path(out).exists()

    wb = load_workbook(out)
    assert wb.sheetnames == ["게시글", "_errors"]

    ws = wb["게시글"]
    assert ws.cell(1, 1).value == "번호"
    assert ws.cell(2, 2).value == "1002"
    assert ws.cell(2, 5).value == "2025-09-12 10:30:00"
    assert ws.cell(2, 7).value == "u1\nu2"          # 이미지 URL
    assert ws.cell(2, 8).value == "p1\np2"          # 이미지 로컬경로

    es = wb["_errors"]
    assert es.cell(1, 1).value == "게시글ID"
    assert es.cell(2, 3).value == "article"
